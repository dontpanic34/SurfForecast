#!/usr/bin/env python3
"""Exporte les traductions de l'appli en fichiers Excel (un par langue) pour une relecture humaine.

Colonnes : N°, Écran, Français, Traduction actuelle, Correction (à remplir), Commentaire.
Usage : export_xlsx.py [dossier de sortie]   (langues : en, de, nl, es)
"""
import glob, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "shared", "src", "commonMain", "kotlin", "com", "surfcast", "surfforecast")
LANGS = [("en", "Anglais"), ("de", "Allemand"), ("nl", "Néerlandais"), ("es", "Espagnol")]
SCREENS = [  # (début du nom de fichier, écran)
    ("NowTide", "Actuellement"), ("ObservedNow", "Actuellement"), ("DailyTimeline", "Déroulé de la journée"), ("ModelVoices", "Modèles partagés"),
    ("WindCard", "Vent"), ("WeatherCard", "Météo"), ("SurfCard", "Houle"), ("HourlyForecast", "Prévisions heure par heure"),
    ("Preferences", "Réglages"), ("Settings", "Réglages"), ("Quiver", "Matériel"), ("Gear", "Matériel"), ("BoardVolume", "Matériel"),
    ("SessionLog", "Journal de session"), ("SessionFeel", "Journal de session"), ("SessionSummary", "Journal de session"),
    ("Comprendre", "Comprendre"), ("LevelQuestionnaire", "Questionnaire de niveau"), ("SurferProfile", "Profil"),
    ("CoachMarks", "Visite guidée"), ("SpotSelection", "Choix du spot"), ("SpotHeader", "Choix du spot"), ("CardInfo", "Réglages d'un encart"),
    ("MainScreen", "Écran principal"), ("SurfController", "Messages de l'appli"), ("SurfRepository", "Messages de l'appli"),
    ("SurfSpotsData", "Noms de spots et régions"), ("DayPhrase", "Phrases du jour"), ("SurfScoring", "Notes et explications"),
    ("ProfileWizard", "Profil"), ("ForecastModels", "Modèles de prévision"), ("ForecastSources", "Sources et modèles"), ("TextDrawing", "Graphiques"),
    ("ForecastHistory", "Historique des relevés"), ("TimeRuler", "Barre de temps"), ("CellBars", "Barre de temps"),
]


def unescape(s):
    return (s.replace('\\n', '\n').replace('\\t', '\t').replace('\\"', '"').replace('\\$', '$').replace('\\\\', '\\'))


STR = r'"((?:[^"\\]|\\.)*)"'
ENTRY = re.compile(r'^\s*' + STR + r'\s+to\s+(?:t\(|listOf\()\s*' + STR + r'\s*,\s*' + STR + r'\s*,\s*' + STR + r'\s*,\s*' + STR, re.M | re.S)


def load_entries():
    out = []
    for f in sorted(glob.glob(os.path.join(SRC, "Translations*.kt"))) + [os.path.join(SRC, "I18n.kt")]:
        for m in ENTRY.finditer(open(f, encoding="utf-8").read()):
            out.append([unescape(g) for g in m.groups()])
    seen, res = set(), []
    for e in out:  # le dernier gagne (CORE de I18n.kt passe après les fichiers générés)
        pass
    d = {}
    for e in out:
        d[e[0]] = e[1:]
    return d


def screen_of(fr, sources):
    for name, text in sources:
        if fr in text:
            for prefix, screen in SCREENS:
                if name.startswith(prefix):
                    return screen
            return name.replace(".kt", "")
    return "Texte construit ou divers"


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    os.makedirs(out, exist_ok=True)
    entries = load_entries()
    sources = []
    for f in sorted(glob.glob(os.path.join(SRC, "*.kt"))):
        n = os.path.basename(f)
        if n.startswith("Translations") or n == "I18n.kt":
            continue
        sources.append((n, open(f, encoding="utf-8").read()))
    rows = []
    for fr, tr in entries.items():
        esc = fr.replace('\\', '\\\\').replace('"', '\\"').replace('\n', '\\n')
        rows.append((screen_of(esc, sources), fr, tr))
    rows.sort(key=lambda r: (r[0], len(r[1]) > 160, r[1].lower()))
    for i, (code, label) in enumerate(LANGS):
        wb = Workbook()
        ws = wb.active
        ws.title = label
        head = ["N°", "Écran", "Français", f"Traduction actuelle ({label})", "Correction (à remplir)", "Commentaire"]
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for n, (screen, fr, tr) in enumerate(rows, 1):
            ws.append([n, screen, fr, tr[i], "", ""])
        for col, w in zip("ABCDEF", (6, 24, 60, 60, 50, 30)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
            row[4].fill = PatternFill("solid", fgColor="FFF2CC")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:F{len(rows) + 1}"
        path = os.path.join(out, f"traduction-{code}.xlsx")
        wb.save(path)
        print(path, len(rows), "lignes")


main()
